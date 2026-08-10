import io

import openpyxl
from docx import Document

import collector_maroc_pps as pps


def test_find_keyword_snippet_trouve():
    text = "Ce document decrit le developpement logiciel prevu pour 2026 dans le cadre du projet."
    snippet = pps._find_keyword_snippet(text, ["developpement logiciel"])
    assert snippet is not None
    assert "developpement logiciel" in snippet.lower()


def test_find_keyword_snippet_absent():
    text = "Ce document decrit des travaux de voirie prevus pour 2026."
    assert pps._find_keyword_snippet(text, ["developpement logiciel", "informatique"]) is None


def test_href_id_org_regex():
    href = "javascript:popUp('index.php?page=commun.PopupListePPsDownloadFile&amp;id=33906&amp;org=g3h');"
    m = pps._HREF_ID_ORG_RE.search(href)
    assert m is not None
    assert m.groups() == ("33906", "g3h")


def test_clean_date():
    assert pps._clean_date("03/08/2026 13:55") == "2026-08-03"
    assert pps._clean_date(None) is None
    assert pps._clean_date("non précisé") is None


def _build_xlsx_bytes(cell_values: list[str]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, value in enumerate(cell_values, start=1):
        ws.cell(row=i, column=1, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_text_from_xlsx_reel():
    """Cas réel constaté : certains documents PPS sont de vrais .xlsx
    téléversés tels quels, pas des scans — auparavant rejetés à tort faute
    de support de ce format."""
    content = _build_xlsx_bytes(["Programme prévisionnel 2026", "Développement d'un système d'information"])
    text = pps._extract_text_from_xlsx(content)
    assert "Développement d'un système d'information" in text


def test_extract_text_from_xlsx_contenu_invalide_retourne_vide():
    assert pps._extract_text_from_xlsx(b"pas un vrai xlsx") == ""


def test_extract_text_from_docx_reel():
    content = _build_docx_bytes(["Programme prévisionnel 2026", "Développement d'une application mobile"])
    text = pps._extract_text_from_docx(content)
    assert "Développement d'une application mobile" in text


def test_extract_text_from_docx_contenu_invalide_retourne_vide():
    assert pps._extract_text_from_docx(b"pas un vrai docx") == ""


class _FakeResponse:
    def __init__(self, body: bytes, ok: bool = True):
        self.ok = ok
        self.content = body


def test_download_and_extract_text_xlsx(monkeypatch):
    xlsx_bytes = _build_xlsx_bytes(["Fourniture de licences Business Intelligence"])
    monkeypatch.setattr(pps.requests, "get", lambda *a, **k: _FakeResponse(xlsx_bytes))
    text = pps.download_and_extract_text("123", "abc")
    assert "Business Intelligence" in text


def test_download_and_extract_text_docx(monkeypatch):
    docx_bytes = _build_docx_bytes(["Mission d'AMOA pour le système d'information"])
    monkeypatch.setattr(pps.requests, "get", lambda *a, **k: _FakeResponse(docx_bytes))
    text = pps.download_and_extract_text("123", "abc")
    assert "AMOA" in text


def test_download_and_extract_text_zip_generique_retourne_vide(monkeypatch):
    """Un zip qui n'est ni xlsx ni docx (constaté en direct) doit retomber
    sur une chaîne vide, jamais une exception."""
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "contenu quelconque")

    monkeypatch.setattr(pps.requests, "get", lambda *a, **k: _FakeResponse(buf.getvalue()))
    assert pps.download_and_extract_text("123", "abc") == ""
